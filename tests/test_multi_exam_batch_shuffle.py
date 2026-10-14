# -*- coding: utf-8 -*-
"""
tests/test_multi_exam_batch_shuffle.py — Kiểm tra tính năng nạp nhiều đề thi cùng lúc,
xóa đề hiện tại, làm mới toàn bộ, trộn nhiều bộ đề hàng loạt và xuất file theo từng thư mục bộ đề.
"""
import os
import sys
import shutil
import pytest
import tkinter as tk
import openpyxl

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from omr_exam.models import ExamSection, ExamQuestion, OptionItem
from omr_exam.exam_shuffler import shuffle_exam
from omr_exam.excel_exporter import export_consolidated_master_excel
from ui.tab_shuffle import TabShuffle, LoadedExamItem
from tkinter import messagebox


@pytest.fixture(autouse=True)
def mock_messageboxes(monkeypatch):
    """Tự động mock tất cả hộp thoại messagebox để tránh treo test."""
    monkeypatch.setattr(messagebox, "showinfo", lambda *args, **kwargs: True)
    monkeypatch.setattr(messagebox, "showwarning", lambda *args, **kwargs: True)
    monkeypatch.setattr(messagebox, "showerror", lambda *args, **kwargs: True)
    monkeypatch.setattr(messagebox, "askyesno", lambda *args, **kwargs: True)


@pytest.fixture
def tk_root():
    """Tạo Tk root window ẩn cho unit test."""
    try:
        root = tk.Tk()
        root.withdraw()
        yield root
        root.destroy()
    except Exception:
        yield None


def _make_sample_section(prefix: str, count: int = 5) -> ExamSection:
    """Tạo 1 phần đề mẫu với số lượng câu hỏi nhất định."""
    questions = []
    for i in range(1, count + 1):
        q = ExamQuestion(
            index=i,
            section="Phần I: Trắc nghiệm",
            qtype="TN",
            content=f"[{prefix}] Câu hỏi số {i}: Giá trị là bao nhiêu?",
            options=[
                OptionItem("A", f"Đáp án A cho câu {i}", is_correct=(i % 4 == 1)),
                OptionItem("B", f"Đáp án B cho câu {i}", is_correct=(i % 4 == 2)),
                OptionItem("C", f"Đáp án C cho câu {i}", is_correct=(i % 4 == 3)),
                OptionItem("D", f"Đáp án D cho câu {i}", is_correct=(i % 4 == 0)),
            ],
            correct_label=["D", "A", "B", "C"][i % 4],
        )
        questions.append(q)
    return ExamSection(title="Phần I: Trắc nghiệm", questions=questions)


def test_loaded_exam_item_properties():
    """Kiểm tra thuộc tính cơ bản của LoadedExamItem."""
    sec = _make_sample_section("Toán", count=10)
    item = LoadedExamItem(name="De_Toan_12.docx", file_path="C:/dummy/De_Toan_12.docx", sections=[sec])
    assert item.name == "De_Toan_12.docx"
    assert item.total_questions == 10
    assert item.tn_count == 10
    assert item.tl_count == 0
    assert item.shuffled_exams == []


def test_tab_shuffle_multi_exam_management(tk_root):
    """Kiểm tra quản lý danh sách nhiều đề, chọn đề, xóa đề và làm mới toàn bộ."""
    if not tk_root:
        pytest.skip("Tkinter display not available")

    tab = TabShuffle(tk_root)

    # 1. Ban đầu rỗng
    assert len(tab._loaded_exams) == 0
    assert tab._current_exam is None

    # 2. Thêm 3 đề thi
    sec1 = _make_sample_section("Toan", 5)
    sec2 = _make_sample_section("Ly", 8)
    sec3 = _make_sample_section("Hoa", 6)

    item1 = LoadedExamItem("De_Toan.docx", "path/to/De_Toan.docx", [sec1])
    item2 = LoadedExamItem("De_Ly.docx", "path/to/De_Ly.docx", [sec2])
    item3 = LoadedExamItem("De_Hoa.docx", "path/to/De_Hoa.docx", [sec3])

    tab._loaded_exams.extend([item1, item2, item3])
    tab._active_exam_idx = 0
    tab._update_exam_list_ui()

    assert len(tab._loaded_exams) == 3
    assert tab._current_exam.name == "De_Toan.docx"
    assert tab._cb_exam_selector["values"] is not None
    assert len(tab._cb_exam_selector["values"]) == 3

    # 3. Đổi sang đề thứ 2
    tab._active_exam_idx = 1
    tab._update_exam_list_ui()
    assert tab._current_exam.name == "De_Ly.docx"
    assert tab._current_exam.total_questions == 8

    # 4. Xóa đề thứ 2 (De_Ly.docx) - gọi logic xóa trực tiếp không pop modal
    tab._loaded_exams.pop(1)
    tab._active_exam_idx = min(tab._active_exam_idx, len(tab._loaded_exams) - 1)
    tab._update_exam_list_ui()

    assert len(tab._loaded_exams) == 2
    assert tab._current_exam.name == "De_Hoa.docx"

    # 5. Làm mới toàn bộ (clear all)
    tab._on_clear_all(silent=True)
    assert len(tab._loaded_exams) == 0
    assert tab._active_exam_idx == -1
    assert tab._current_exam is None
    assert tab._lbl_exam_count.cget("text") == "(Chưa có đề nào)"


def test_batch_shuffle_multiple_exams(tk_root):
    """Kiểm tra xáo trộn đồng thời nhiều bộ đề thành các tập mã đề độc lập."""
    if not tk_root:
        pytest.skip("Tkinter display not available")

    tab = TabShuffle(tk_root)

    sec1 = _make_sample_section("Toán 12", 4)
    sec2 = _make_sample_section("Vật Lý 12", 4)

    item1 = LoadedExamItem("De_Toan_12.docx", None, [sec1])
    item2 = LoadedExamItem("De_VatLy_12.docx", None, [sec2])
    tab._loaded_exams = [item1, item2]
    tab._active_exam_idx = 0
    tab._update_exam_list_ui()

    tab._var_batch_scope.set("all")
    tab._spin_codes.set("3")
    tab._entry_start_code.delete(0, "end")
    tab._entry_start_code.insert(0, "201")
    tab._entry_custom_codes.delete(0, "end")

    # Thực hiện trộn batch
    tab._on_do_shuffle()

    # Kiểm tra cả 2 bộ đề đều đã được trộn
    assert len(item1.shuffled_exams) == 3
    assert len(item2.shuffled_exams) == 3
    assert set(item1.keys_by_code.keys()) == {"201", "202", "203"}
    assert set(item2.keys_by_code.keys()) == {"201", "202", "203"}

    # Kiểm tra ma trận hoán vị có đầy đủ
    assert len(item1.question_map_by_code["201"]) == 4
    assert len(item2.question_map_by_code["201"]) == 4

    # Kiểm tra Card 3 combobox có 2 bộ kết quả
    assert len(tab._cb_result_exam_selector["values"]) == 2


def test_consolidated_master_excel_generation(tmp_path):
    """Kiểm tra tạo file Excel tổng hợp đáp án cho nhiều bộ đề thi."""
    excel_file = tmp_path / "Tong_Hop_Dap_An.xlsx"

    sec1 = _make_sample_section("Toan", 5)
    sec2 = _make_sample_section("Ly", 5)

    shuffled1 = shuffle_exam([sec1], custom_codes=["101", "102"])
    shuffled2 = shuffle_exam([sec2], custom_codes=["201", "202"])

    datasets = [
        {
            "name": "De_Thi_Toan_12",
            "keys_by_code": {ex.code: ex.answer_key for ex in shuffled1},
            "question_map_by_code": {ex.code: ex.question_map for ex in shuffled1},
            "option_map_by_code": {ex.code: getattr(ex, "option_map", {}) for ex in shuffled1},
        },
        {
            "name": "De_Thi_VatLy_12",
            "keys_by_code": {ex.code: ex.answer_key for ex in shuffled2},
            "question_map_by_code": {ex.code: ex.question_map for ex in shuffled2},
            "option_map_by_code": {ex.code: getattr(ex, "option_map", {}) for ex in shuffled2},
        }
    ]

    res_path = export_consolidated_master_excel(str(excel_file), datasets)
    assert os.path.exists(res_path)

    wb = openpyxl.load_workbook(res_path)
    sheet_names = wb.sheetnames

    # Kiểm tra trang tổng quan
    assert "Tổng Quan Các Bộ Đề" in sheet_names
    ws_main = wb["Tổng Quan Các Bộ Đề"]
    assert ws_main.cell(row=1, column=1).value == "BẢNG TỔNG HỢP CÁC BỘ ĐỀ THI ĐÃ TRỘN"
    # Có 2 dòng dữ liệu ở row 4, row 5
    assert ws_main.cell(row=4, column=2).value == "De_Thi_Toan_12"
    assert ws_main.cell(row=5, column=2).value == "De_Thi_VatLy_12"

    # Kiểm tra các trang đáp án của từng bộ
    assert "Bo_01_DapAn" in sheet_names
    assert "Bo_01_SoiCau" in sheet_names
    assert "Bo_02_DapAn" in sheet_names
    assert "Bo_02_SoiCau" in sheet_names

    ws_k1 = wb["Bo_01_DapAn"]
    assert "101" in str(ws_k1.cell(row=3, column=2).value)
    assert "102" in str(ws_k1.cell(row=3, column=3).value)

    wb.close()


def test_multi_exam_export_folder_structure(tk_root, tmp_path, monkeypatch):
    """Kiểm tra xuất toàn bộ file Word & Excel cho nhiều bộ đề vào từng thư mục riêng biệt."""
    if not tk_root:
        pytest.skip("Tkinter display not available")

    # Mock askdirectory và startfile
    out_dir = tmp_path / "xuat_nhieu_bo_de"
    out_dir.mkdir(parents=True, exist_ok=True)
    monkeypatch.setattr("tkinter.filedialog.askdirectory", lambda *a, **k: str(out_dir))
    monkeypatch.setattr(os, "startfile", lambda *a, **k: None)

    tab = TabShuffle(tk_root)
    sec1 = _make_sample_section("Toán", 3)
    sec2 = _make_sample_section("Vật Lý", 3)

    item1 = LoadedExamItem("De_Thi_Toan_12.docx", None, [sec1])
    item2 = LoadedExamItem("De_Thi_VatLy_12.docx", None, [sec2])
    tab._loaded_exams = [item1, item2]
    tab._active_exam_idx = 0
    tab._update_exam_list_ui()

    tab._var_batch_scope.set("all")
    tab._spin_codes.set("2")
    tab._entry_start_code.delete(0, "end")
    tab._entry_start_code.insert(0, "101")
    tab._entry_custom_codes.delete(0, "end")

    tab._on_do_shuffle()

    # Xuất file
    tab._on_export_docx_files()

    # Kiểm tra thư mục con đã được tạo: "Bộ 01 - De_Thi_Toan_12" và "Bộ 02 - De_Thi_VatLy_12"
    subdirs = [d.name for d in out_dir.iterdir() if d.is_dir()]
    assert any("Bộ 01" in s for s in subdirs)
    assert any("Bộ 02" in s for s in subdirs)

    # Kiểm tra bên trong mỗi thư mục con có file .docx
    for s_name in subdirs:
        s_path = out_dir / s_name
        docx_files = list(s_path.glob("*.docx"))
        assert len(docx_files) >= 2

    # Kiểm tra file Excel tổng hợp tại thư mục gốc
    master_excel = out_dir / "Tong_Hop_Dap_An_Va_Ma_Tran_Tat_Ca_Bo_De.xlsx"
    assert master_excel.exists()

