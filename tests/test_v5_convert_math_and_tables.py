import pytest
from ui.tab_convert import (
    normalize_math_to_braced_dollars,
    TabConvert,
)


def test_normalize_math_to_braced_dollars():
    # 1. Math formulas with single dollar
    s1 = "Cho hàm số $y = f(x)$ liên tục trên $\\mathbb{R}$."
    res1 = normalize_math_to_braced_dollars(s1)
    assert res1 == "Cho hàm số ${y = f(x)}$ liên tục trên ${\\mathbb{R}}$."

    # 2. Math formulas already having braces
    s2 = "Biết ${\\int_0^1 f(x)dx = 3}$ và $I = 5$."
    res2 = normalize_math_to_braced_dollars(s2)
    assert res2 == "Biết ${\\int_0^1 f(x)dx = 3}$ và ${I = 5}$."
    assert "${{I = 5}}$" not in res2
    assert "${{\\int_0^1" not in res2

    # 3. Double dollar block math
    s3 = "Tính: $$\\lim_{x \\to 0} \\frac{\\sin x}{x} = 1$$"
    res3 = normalize_math_to_braced_dollars(s3)
    assert res3 == "Tính: ${\\lim_{x \\to 0} \\frac{\\sin x}{x} = 1}$"

    # 4. LaTeX brackets \( ... \) và \[ ... \]
    s4 = "Góc \\(\\widehat{ABC} = 60^\\circ\\) và biểu thức \\[E = mc^2\\]"
    res4 = normalize_math_to_braced_dollars(s4)
    assert res4 == "Góc ${\\widehat{ABC} = 60^\\circ}$ và biểu thức ${E = mc^2}$"


def test_table_line_detection_and_math_pipe_masking():
    # Bảng bình thường
    t1 = "| Cột 1 | Cột 2 | Cột 3 |"
    assert TabConvert._is_table_line(t1) is True

    # Bảng phân tách bởi |
    t2 = "Cột 1 | Cột 2 | Cột 3"
    assert TabConvert._is_table_line(t2) is True

    # Bảng có chứa công thức giá trị tuyệt đối |x|
    t3 = "| Hàm số | Điều kiện $|x| < 3$ | Ghi chú |"
    assert TabConvert._is_table_line(t3) is True

    # Dòng văn bản bình thường chứa |x| KHÔNG phải bảng
    t4 = "Phương trình có nghiệm khi $|x| < 5$ hoặc $|x| > 10$."
    assert TabConvert._is_table_line(t4) is False

    # Dòng bảng biến thiên
    t5 = "| x | -\\infty | 0 | +\\infty |"
    t6 = "| y' | + | 0 | - |"
    assert TabConvert._is_table_line(t5) is True
    assert TabConvert._is_table_line(t6) is True


def test_parse_table_row_preserves_math_pipes():
    row_text = "| Hàm số | Điều kiện $|x| < 3$ | Ghi chú |"
    cells = TabConvert._parse_table_row(row_text)
    assert len(cells) == 3
    assert cells[0] == "Hàm số"
    # Dấu pipe bên trong công thức không bị cắt làm đôi và được chuẩn hóa ${...}$
    assert cells[1] == "Điều kiện ${|x| < 3}$"
    assert cells[2] == "Ghi chú"


def test_build_word_from_pages_creates_table_and_preserves_math(tmp_path):
    from docx import Document
    out_docx = str(tmp_path / "test_convert.docx")

    raw_text = (
        "Câu 1: Cho hàm số $y = x^2 - 4x + 3$. Bảng biến thiên:\n"
        "| x | -\\infty | 2 | +\\infty |\n"
        "|---|---|---|---|\n"
        "| y' | - | 0 | + |\n"
        "| y | +\\infty | -1 | +\\infty |\n\n"
        "Phương trình có 2 nghiệm phân biệt khi $m > -1$."
    )

    page_results = [{
        "text": raw_text,
        "images": [],
        "page_num": 1,
    }]

    # Tạo một instance giả của TabConvert để gọi _build_word_from_pages
    converter = TabConvert.__new__(TabConvert)
    converter._build_word_from_pages(page_results, out_docx)

    # Đọc lại docx đã tạo để xác thực
    doc = Document(out_docx)

    # 1. Kiểm tra bảng được tạo thành công
    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert len(table.rows) >= 3
    assert len(table.columns) == 4

    # 2. Kiểm tra các paragraph có công thức toán được bọc trong ${...}$
    text_content = "\n".join([p.text for p in doc.paragraphs])
    assert "${y = x^2 - 4x + 3}$" in text_content
    assert "${m > -1}$" in text_content
    assert "$y = x^2 - 4x + 3$" not in text_content # Không còn dạng trần thiếu {}
