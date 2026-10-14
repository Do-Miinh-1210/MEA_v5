"""
tests/test_parser.py — Test core/parser.py.

Coverage:
  • Parse user's real matrix file (regression for the "1 dòng rác" bug)
  • Parse synthetic Simple 6-col matrix
  • Code extraction (case-insensitive, TLN, lowercase c)
  • Level normalization (strict vs partial match)
  • DS grouping by code-number prefix
"""
from __future__ import annotations

import pytest

from core.parser import (
    parse_matrix_file,
    _normalize_level,
    _strict_level_header,
    _extract_codes_from_cell,
    _detect_qtype,
)


# ══════════════════════════════════════════════════════════════════════════
# REGRESSION: user's real matrix file → must produce 24 rows (16 TN + 4 DS + 4 TLN)
# ══════════════════════════════════════════════════════════════════════════

@pytest.mark.docx
def test_user_matrix_regression(user_matrix_file):
    """Sửa bug `_parse_level_col_spec_table` ăn nhầm → vẫn không lặp lại."""
    result = parse_matrix_file(str(user_matrix_file))

    # Tổng câu — đúng
    assert len(result.rows) == 24, (
        f"Expected 24 rows, got {len(result.rows)}. "
        f"Bug `_parse_level_col_spec_table` đã quay lại?"
    )

    # Phân loại
    tn = sum(1 for r in result.rows if r.qtype == "TN")
    ds = sum(1 for r in result.rows if r.qtype == "DS")
    tln = sum(1 for r in result.rows if r.qtype == "TLN")
    tl = sum(1 for r in result.rows if r.qtype == "TL")

    assert tn == 16, f"Expected 16 TN, got {tn}"
    assert ds == 4, f"Expected 4 DS (đã group abcd), got {ds}"
    assert tln == 4, f"Expected 4 TLN (TLN1-TLN4), got {tln}"
    assert tl == 0, f"Expected 0 pure TL, got {tl}"

    # Tổng điểm = 10đ
    total = sum(r.points for r in result.rows)
    assert abs(total - 10.0) < 0.01, f"Expected 10.0đ, got {total:.2f}đ"

    # DS phải có 4 ý a/b/c/d
    for r in result.rows:
        if r.qtype == "DS":
            assert r.ds_items, f"DS row '{r.codes}' missing ds_items"
            labels = {it.label for it in r.ds_items}
            assert labels == {"a", "b", "c", "d"}, (
                f"DS row codes={r.codes} should have labels a/b/c/d, got {labels}"
            )


# ══════════════════════════════════════════════════════════════════════════
# UNIT: _normalize_level vs _strict_level_header
# ══════════════════════════════════════════════════════════════════════════

class TestLevelNormalize:
    """`_normalize_level` permissive (cho phép partial match);
    `_strict_level_header` strict (chỉ short label) — dùng cho header detection."""

    def test_exact_codes(self):
        assert _normalize_level("NB") == "NB"
        assert _normalize_level("TH") == "TH"
        assert _normalize_level("VD") == "VD"
        assert _normalize_level("VDC") == "VD"   # VDC → VD

    def test_full_names(self):
        assert _normalize_level("Nhận biết") == "NB"
        assert _normalize_level("Thông hiểu") == "TH"
        assert _normalize_level("Vận dụng") == "VD"

    def test_partial_match(self):
        # Permissive — match anywhere
        assert _normalize_level("Nhận biết khái niệm nguyên hàm") == "NB"

    def test_strict_rejects_long(self):
        # Strict — chỉ chấp nhận label NGẮN (header cell)
        assert _strict_level_header("Nhận biết khái niệm nguyên hàm") == ""
        assert _strict_level_header("NB") == "NB"
        assert _strict_level_header("Nhận biết") == "NB"
        # Quá dài (>25 char) → reject
        assert _strict_level_header("Nhận biết được khái niệm và công thức") == ""


# ══════════════════════════════════════════════════════════════════════════
# UNIT: _extract_codes_from_cell (case-insensitive, TLN, lowercase c)
# ══════════════════════════════════════════════════════════════════════════

class TestExtractCodes:
    def test_uppercase_C(self):
        assert _extract_codes_from_cell("C1") == ["C1"]
        assert _extract_codes_from_cell("C13a") == ["C13a"]

    def test_lowercase_c(self):
        # Bug cũ: chỉ match C hoa → c1 bị bỏ
        assert _extract_codes_from_cell("c1") == ["C1"]
        assert _extract_codes_from_cell("c10") == ["C10"]
        assert _extract_codes_from_cell("c12, \nc13") == ["C12", "C13"]

    def test_TLN_codes(self):
        # Bug cũ: TLN không match (TL\d+ requires digit ngay sau TL, chữ N chen vào)
        assert _extract_codes_from_cell("TLN1") == ["TLN1"]
        assert _extract_codes_from_cell("TLN2") == ["TLN2"]
        assert _extract_codes_from_cell("TLN10") == ["TLN10"]

    def test_DS_codes(self):
        # 1a, 2b, 3c — DS items
        assert _extract_codes_from_cell("1a") == ["1a"]
        assert _extract_codes_from_cell("2b, \n2c") == ["2b", "2c"]

    def test_cau_format(self):
        assert _extract_codes_from_cell("Câu 1") == ["Câu 1"]
        assert _extract_codes_from_cell("Câu 1, 2, 3") == [
            "Câu 1", "Câu 2", "Câu 3",
        ]

    def test_NL_annotation_stripped(self):
        # `(NL 1.1)` annotations should be removed before code extraction
        assert _extract_codes_from_cell("C1 (NL 1.1)") == ["C1"]

    def test_no_codes(self):
        assert _extract_codes_from_cell("") == []
        assert _extract_codes_from_cell("Nhận biết khái niệm") == []


# ══════════════════════════════════════════════════════════════════════════
# UNIT: _detect_qtype (TLN must come before TL)
# ══════════════════════════════════════════════════════════════════════════

class TestDetectQtype:
    def test_explicit_dtype(self):
        assert _detect_qtype("NPA", "") == "TN"
        assert _detect_qtype("Đ-S", "") == "DS"
        assert _detect_qtype("ĐS", "") == "DS"
        assert _detect_qtype("TLN", "") == "TLN"
        assert _detect_qtype("Tự luận", "") == "TL"

    def test_infer_from_codes(self):
        # Empty dtype → suy từ code
        assert _detect_qtype("", "C1") == "TN"
        assert _detect_qtype("", "TL1") == "TL"
        assert _detect_qtype("", "TLN1") == "TLN"   # Bug cũ: TLN1 → TL
        assert _detect_qtype("", "1a") == "DS"


# ══════════════════════════════════════════════════════════════════════════
# DOCX-end-to-end (synthetic — không cần file ngoài)
# ══════════════════════════════════════════════════════════════════════════

@pytest.fixture
def synthetic_simple_matrix(tmp_path):
    """Tạo file .docx Simple 6-col tối giản trong tmp_path."""
    try:
        from docx import Document
    except ImportError:
        pytest.skip("python-docx not installed")

    doc = Document()
    doc.add_paragraph("MA TRẬN ĐỀ KIỂM TRA — Toán 12 (Test)")
    doc.add_paragraph("(Thời gian làm bài 90 phút)")
    doc.add_paragraph("BẢNG ĐẶC TẢ")

    # Bảng spec 6 cột
    tbl = doc.add_table(rows=4, cols=6)
    hdr = tbl.rows[0].cells
    hdr[0].text = "STT"
    hdr[1].text = "Kiến thức"
    hdr[2].text = "Mức độ"
    hdr[3].text = "Dạng"
    hdr[4].text = "Câu"
    hdr[5].text = "Yêu cầu cần đạt"

    rows = [
        ("1", "Nguyên hàm", "NB", "NPA", "C1", "Nhận biết nguyên hàm"),
        ("2", "Tích phân", "TH", "NPA", "C2", "Tính tích phân cơ bản"),
        ("3", "Tích phân", "VD", "Đ-S", "1a", "Đúng/sai tích phân từng phần"),
    ]
    for i, (stt, topic, lvl, dtype, code, req) in enumerate(rows, start=1):
        for j, val in enumerate([stt, topic, lvl, dtype, code, req]):
            tbl.rows[i].cells[j].text = val

    out = tmp_path / "synthetic.docx"
    doc.save(out)
    return out


def test_parse_synthetic_simple(synthetic_simple_matrix):
    """Parser phải đọc đúng Simple 6-col."""
    result = parse_matrix_file(str(synthetic_simple_matrix))
    assert len(result.rows) >= 3
    # Có 2 TN + 1 DS
    qtypes = [r.qtype for r in result.rows]
    assert "TN" in qtypes
    assert "DS" in qtypes


def test_parse_returns_title(synthetic_simple_matrix):
    """Title phải được trích từ paragraph chứa 'MA TRẬN'."""
    result = parse_matrix_file(str(synthetic_simple_matrix))
    assert "MA TRẬN" in result.title.upper()


def test_parse_returns_time_str(synthetic_simple_matrix):
    """Time '90 phút' phải được trích."""
    result = parse_matrix_file(str(synthetic_simple_matrix))
    assert "90" in result.time_str
