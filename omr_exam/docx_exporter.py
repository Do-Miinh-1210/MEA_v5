"""
omr_exam/docx_exporter.py — Xuất các mã đề thi đã trộn ra file Word (.docx) & Excel (.xlsx)
Đảm bảo:
- Bố cục đề thi rõ ràng, tiêu đề mã đề.
- Hiển thị A, B, C, D đều đặn.
- Nếu có câu Tự luận (TL), tạo các dòng kẻ chấm chấm để học sinh viết đáp án.
- Xuất từng file Word riêng lẻ + File Word GỘP tất cả các mã đề.
- Xuất Bảng đáp án Word (.docx) & Bảng đáp án + Ma trận hoán vị Excel (.xlsx).
"""
from __future__ import annotations
import os
import copy
from pathlib import Path
from typing import List, Dict, Optional, Set
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from omr_exam.models import ShuffledExamCode, ExamQuestion
from omr_exam.excel_exporter import export_exam_keys_and_matrix_to_excel


def _add_math_run_or_text(paragraph, text: str, font_size: int = 12, bold: bool = False, italic: bool = False):
    """Thêm text vào paragraph Word ở dạng văn bản chuẩn Times New Roman.
    Tuyệt đối KHÔNG tự động chuyển sang Word Equation (OMML).
    Giữ nguyên công thức $...$ để tương thích hoàn toàn với MathType (Toggle TeX / Convert Equations) trên Word.
    """
    if "$" in text:
        try:
            from export.math_converter import split_text_and_math
            parts = split_text_and_math(text)
            for is_math, content in parts:
                if is_math:
                    r = paragraph.add_run(f"${content}$")
                    r.font.name = "Times New Roman"
                    r.font.size = Pt(font_size)
                    r.font.italic = True
                    r.font.bold = bold
                elif content:
                    r = paragraph.add_run(content)
                    r.font.name = "Times New Roman"
                    r.font.size = Pt(font_size)
                    r.font.bold = bold
                    r.font.italic = italic
            return
        except Exception:
            pass
    r = paragraph.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(font_size)
    r.font.bold = bold
    r.font.italic = italic


def export_shuffled_exams_to_docx(
    shuffled_exams: List[ShuffledExamCode],
    output_dir: str,
    school_name: str = "SỞ GD&ĐT ....................",
    exam_title: str = "BÀI KIỂM TRA ĐỊNH KỲ",
    subject_name: str = "MÔN: TOÁN HỌC",
    time_minutes: int = 45,
    source_docx_path: Optional[str] = None,
    shuffle_mode: str = "both",
    auto_pin_dependent_options: bool = True,
    pinned_question_indices: Optional[Set[int]] = None,
    export_merged_docx: bool = True,
    is_ai_variant: bool = False,
) -> List[str]:
    """
    Xuất danh sách các mã đề ra từng file .docx riêng biệt + file gộp + file Excel.
    Nếu source_docx_path được cung cấp và tồn tại:
    Sử dụng công nghệ OpenXML trực tiếp để bảo toàn 100% công thức MathType, OMML,
    hình ảnh, cấu trúc bảng và định dạng của đề gốc.
    """
    out_path = Path(output_dir)
    out_path.mkdir(parents=True, exist_ok=True)
    
    if source_docx_path and os.path.exists(source_docx_path):
        from omr_exam.docx_shuffler_xml import shuffle_docx_with_full_xml_preservation
        num_codes = len(shuffled_exams) if shuffled_exams else 4
        start_code = 101
        custom_codes = [str(ex.code) for ex in shuffled_exams] if shuffled_exams else None
        
        files, keys_by_code, ans_file, excel_file, q_map, opt_map = shuffle_docx_with_full_xml_preservation(
            source_docx_path=source_docx_path,
            output_dir=output_dir,
            num_codes=num_codes,
            start_code=start_code,
            custom_codes=custom_codes,
            shuffle_mode=shuffle_mode,
            auto_pin_dependent_options=auto_pin_dependent_options,
            pinned_question_indices=pinned_question_indices,
            export_merged_docx=export_merged_docx,
            shuffled_exams=shuffled_exams,
            is_ai_variant=is_ai_variant,
        )
        for ex in shuffled_exams:
            if ex.code in keys_by_code:
                ex.answer_key = keys_by_code[ex.code]
            if ex.code in q_map:
                ex.question_map = q_map[ex.code]
            if ex.code in opt_map:
                ex.option_map = opt_map[ex.code]
        return files

    generated_files = []
    
    # Chuẩn bị file gộp nếu được chọn
    merged_doc = Document() if export_merged_docx else None
    if merged_doc:
        for section in merged_doc.sections:
            section.top_margin = Inches(0.7)
            section.bottom_margin = Inches(0.7)
            section.left_margin = Inches(0.8)
            section.right_margin = Inches(0.8)
            
    for exam_idx, exam in enumerate(shuffled_exams):
        filename = f"De_Thi_Ma_{exam.code}.docx"
        filepath = out_path / filename
        
        doc = Document()
        
        # Thiết lập lề trang
        for section in doc.sections:
            section.top_margin = Inches(0.7)
            section.bottom_margin = Inches(0.7)
            section.left_margin = Inches(0.8)
            section.right_margin = Inches(0.8)
            
        def _populate_doc(target_doc, ex):
            # 1. Header đề thi (2 cột: Trường/Lớp & Tên kỳ thi/Mã đề)
            table = target_doc.add_table(rows=2, cols=2)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = False
            
            # Cột trái
            cell_tl = table.cell(0, 0)
            p = cell_tl.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(school_name + "\nTRƯỜNG THPT: ....................")
            r.font.size = Pt(11)
            r.font.name = "Times New Roman"
            
            cell_bl = table.cell(1, 0)
            p = cell_bl.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run("Họ và tên thí sinh: .......................................\nSố báo danh: .................................................")
            r.font.size = Pt(11)
            r.font.name = "Times New Roman"
            
            # Cột phải
            cell_tr = table.cell(0, 1)
            p = cell_tr.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(f"{exam_title}\n{subject_name}\nThời gian: {time_minutes} phút")
            r.font.bold = True
            r.font.size = Pt(11)
            r.font.name = "Times New Roman"
            
            cell_br = table.cell(1, 1)
            p = cell_br.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(f"MÃ ĐỀ THI: {ex.code}")
            r.font.bold = True
            r.font.size = Pt(14)
            r.font.color.rgb = RGBColor(180, 0, 0)
            r.font.name = "Times New Roman"
            
            target_doc.add_paragraph()  # khoảng cách
            
            # 2. Nội dung các phần và câu hỏi
            for sec in ex.sections:
                p_sec = target_doc.add_paragraph()
                r_sec = p_sec.add_run(sec.title.upper())
                r_sec.font.bold = True
                r_sec.font.size = Pt(12)
                r_sec.font.name = "Times New Roman"
                
                for q in sec.questions:
                    p_q = target_doc.add_paragraph()
                    p_q.paragraph_format.space_before = Pt(4)
                    p_q.paragraph_format.space_after = Pt(2)
                    
                    r_num = p_q.add_run(f"Câu {q.index}. ")
                    r_num.font.bold = True
                    r_num.font.name = "Times New Roman"
                    
                    _add_math_run_or_text(p_q, q.content, font_size=12)
                    
                    if q.qtype == "TN" and q.options:
                        opts_text = "        ".join([f"{opt.label}. {opt.text}" for opt in q.options])
                        if len(opts_text) > 85 or any("\n" in opt.text for opt in q.options):
                            for opt in q.options:
                                p_single = target_doc.add_paragraph()
                                p_single.paragraph_format.left_indent = Inches(0.3)
                                p_single.paragraph_format.space_after = Pt(1)
                                r_opt_lbl = p_single.add_run(f"{opt.label}. ")
                                r_opt_lbl.font.bold = True
                                r_opt_lbl.font.name = "Times New Roman"
                                _add_math_run_or_text(p_single, opt.text, font_size=12)
                        else:
                            p_opts = target_doc.add_paragraph()
                            p_opts.paragraph_format.left_indent = Inches(0.25)
                            p_opts.paragraph_format.space_after = Pt(4)
                            for opt in q.options:
                                r_opt_lbl = p_opts.add_run(f"{opt.label}. ")
                                r_opt_lbl.font.bold = True
                                r_opt_lbl.font.name = "Times New Roman"
                                _add_math_run_or_text(p_opts, opt.text, font_size=12)
                                p_opts.add_run("    ")
                    elif q.qtype == "TL":
                        lines_needed = max(q.tl_min_lines, 3) + 2
                        for _ in range(lines_needed):
                            p_line = target_doc.add_paragraph()
                            p_line.paragraph_format.space_after = Pt(2)
                            r_dots = p_line.add_run(". . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . .")
                            r_dots.font.color.rgb = RGBColor(160, 160, 160)
                            r_dots.font.size = Pt(10)
                            
            # 3. Footer kết thúc
            p_end = target_doc.add_paragraph()
            p_end.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r_end = p_end.add_run("---------- HẾT ----------\n(Cán bộ coi thi không giải thích gì thêm)")
            r_end.font.italic = True
            r_end.font.size = Pt(10)
            r_end.font.name = "Times New Roman"

        # Điền vào doc đơn lẻ
        _populate_doc(doc, exam)
        doc.save(str(filepath))
        generated_files.append(str(filepath))
        
        # Nếu có merged doc, điền vào merged doc
        if merged_doc:
            _populate_doc(merged_doc, exam)
            if exam_idx < len(shuffled_exams) - 1:
                merged_doc.add_page_break()
                
    # Lưu file Word GỘP nếu có
    if merged_doc:
        merged_file = out_path / "De_Thi_Tat_Ca_Ma_De_Gop.docx"
        merged_doc.save(str(merged_file))
        generated_files.append(str(merged_file))
        
    # Tạo thêm file Bảng đáp án tổng hợp Word
    ans_key_file = out_path / "Bang_Dap_An_Cac_Ma_De.docx"
    _export_answer_keys_docx(shuffled_exams, str(ans_key_file))
    generated_files.append(str(ans_key_file))
    
    # Tạo file Bảng đáp án & Ma trận hoán vị Excel
    excel_file = out_path / "Bang_Dap_An_Va_Ma_Tran_Hoan_Vi.xlsx"
    keys_by_code = {ex.code: ex.answer_key for ex in shuffled_exams}
    question_map_by_code = {ex.code: ex.question_map for ex in shuffled_exams}
    option_map_by_code = {ex.code: getattr(ex, "option_map", {}) for ex in shuffled_exams}
    
    export_exam_keys_and_matrix_to_excel(
        str(excel_file),
        keys_by_code=keys_by_code,
        question_map_by_code=question_map_by_code,
        orig_keys=None,
        option_map_by_code=option_map_by_code
    )
    generated_files.append(str(excel_file))
    
    return generated_files


def _export_answer_keys_docx(shuffled_exams: List[ShuffledExamCode], filepath: str):
    """Xuất bảng tổng hợp đáp án các mã đề để giáo viên lưu trữ đối chiếu."""
    doc = Document()
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p_title.add_run("BẢNG ĐÁP ÁN CÁC MÃ ĐỀ THI")
    r.font.bold = True
    r.font.size = Pt(16)
    r.font.color.rgb = RGBColor(31, 73, 125)
    r.font.name = "Times New Roman"
    
    max_q = 0
    for ex in shuffled_exams:
        if ex.answer_key:
            max_q = max(max_q, max(ex.answer_key.keys()))
            
    if max_q == 0:
        doc.save(filepath)
        return
        
    table = doc.add_table(rows=max_q + 1, cols=len(shuffled_exams) + 1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    
    # Header bảng
    table.cell(0, 0).paragraphs[0].add_run("Câu").font.bold = True
    for col_idx, ex in enumerate(shuffled_exams):
        table.cell(0, col_idx + 1).paragraphs[0].add_run(f"Đề {ex.code}").font.bold = True
        
    for q_idx in range(1, max_q + 1):
        table.cell(q_idx, 0).paragraphs[0].add_run(str(q_idx))
        for col_idx, ex in enumerate(shuffled_exams):
            ans = ex.answer_key.get(q_idx, "-")
            cell_p = table.cell(q_idx, col_idx + 1).paragraphs[0]
            cell_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r_ans = cell_p.add_run(ans)
            if ans != "-":
                r_ans.font.bold = True
                
    doc.save(filepath)
