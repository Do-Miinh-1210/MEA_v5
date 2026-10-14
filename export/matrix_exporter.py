"""
export/matrix_exporter.py — Xuất ma trận đã chỉnh sửa (CSV UTF-8 / Word bảng đặc tả).
Xuất Word theo đúng chuẩn BGD: nhóm theo chủ đề, cột mức độ × loại câu.
"""

from __future__ import annotations

import csv
import re
from pathlib import Path
from collections import OrderedDict

from core.models import MatrixData, MatrixRow


def export_matrix_csv(matrix: MatrixData, output_path: str) -> None:
    """Xuất CSV (UTF-8 BOM) để mở tốt trong Excel."""
    path = Path(output_path)
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(
            ["STT", "Bật", "Chủ đề", "Loại", "Mức", "Mã câu", "Yêu cầu cần đạt", "Điểm"]
        )
        for r in matrix.rows:
            w.writerow(
                [
                    r.stt,
                    "Có" if r.enabled else "Không",
                    r.topic,
                    r.qtype,
                    r.level,
                    ";".join(r.codes),
                    r.requirement.replace("\n", " ").replace("\r", " "),
                    f"{r.points:g}",
                ]
            )


# ══════════════════════════════════════════════════════════════════════════════
#  WORD EXPORT — Ma trận đặc tả theo chuẩn BGD
# ══════════════════════════════════════════════════════════════════════════════

_QTYPE_ORDER = ["TN", "DS", "TLN", "TL"]
_QTYPE_LABEL = {"TN": "Nhiều\nlựa chọn", "DS": "Đúng –\nSai",
                "TLN": "Trả lời\nngắn", "TL": "Tự luận"}
_LEVEL_ORDER = ["NB", "TH", "VD"]
_LEVEL_LABEL = {"NB": "Nhận biết", "TH": "Thông hiểu", "VD": "Vận dụng"}


def _group_rows_by_topic(rows: list[MatrixRow]) -> OrderedDict:
    """Nhóm rows theo topic, giữ nguyên thứ tự xuất hiện."""
    groups: OrderedDict[str, list[MatrixRow]] = OrderedDict()
    for r in rows:
        topic = r.topic or "(Không rõ)"
        if topic not in groups:
            groups[topic] = []
        groups[topic].append(r)
    return groups


def _detect_qtypes(rows: list[MatrixRow]) -> list[str]:
    """Phát hiện các loại câu có trong ma trận, giữ thứ tự chuẩn."""
    present = set()
    for r in rows:
        present.add(r.qtype)
    return [qt for qt in _QTYPE_ORDER if qt in present]


def _collect_codes(topic_rows: list[MatrixRow], qtype: str, level: str) -> str:
    """Thu thập mã câu cho ô (qtype × level)."""
    codes = []
    for r in topic_rows:
        if r.qtype == qtype and r.level == level and r.enabled:
            codes.extend(r.codes)
            # DS items
            if r.qtype == "DS" and r.ds_items:
                for item in r.ds_items:
                    if item.level == level and item.code and item.code not in codes:
                        codes.append(item.code)
    return ", ".join(codes)


def _count_questions(topic_rows: list[MatrixRow], qtype: str, level: str) -> int:
    """Đếm số câu cho ô (qtype × level)."""
    count = 0
    for r in topic_rows:
        if r.qtype == qtype and r.level == level and r.enabled:
            count += 1
    return count


def _get_requirements(topic_rows: list[MatrixRow]) -> str:
    """Gộp yêu cầu cần đạt cho một chủ đề."""
    reqs = []
    seen = set()
    for r in topic_rows:
        req = r.requirement.strip()
        if req and req not in seen:
            seen.add(req)
            reqs.append(req)
    return "\n".join(reqs)


def _sum_points(topic_rows: list[MatrixRow], qtype: str, level: str) -> float:
    """Tổng điểm cho ô (qtype × level)."""
    total = 0.0
    for r in topic_rows:
        if r.qtype == qtype and r.level == level and r.enabled:
            total += r.points
    return total


def _sum_points_all_qtypes(topic_rows: list[MatrixRow], level: str) -> float:
    """Tổng điểm tất cả qtype cho 1 level."""
    total = 0.0
    for r in topic_rows:
        if r.level == level and r.enabled:
            total += r.points
    return total


def _count_all_qtypes(topic_rows: list[MatrixRow], level: str) -> int:
    """Đếm số câu tất cả qtype cho 1 level."""
    return sum(1 for r in topic_rows if r.level == level and r.enabled)


# ── Column mapping cố định theo chuẩn BGD (19 cột) ──
# Col 0:  TT
# Col 1:  Chủ đề / Chương
# Col 2:  Nội dung / Đơn vị kiến thức
# Col 3-5:  Nhiều lựa chọn (TN) × NB/TH/VD
# Col 6-8:  Đúng–Sai (DS) × NB/TH/VD
# Col 9-11: Trả lời ngắn (TLN) × NB/TH/VD
# Col 12-14: Tự luận (TL) × NB/TH/VD
# Col 15-17: Tổng × NB/TH/VD
# Col 18: Tỉ lệ % điểm số

_COL_QTYPE_START = {"TN": 3, "DS": 6, "TLN": 9, "TL": 12}
_COL_TOTAL_START = 15
_COL_PERCENT = 18
_N_COLS = 19


def export_matrix_docx(matrix: MatrixData, output_path: str) -> None:
    """Xuất Word ma trận đề theo chuẩn BGD Công văn 7991/BGDĐT ngày 17/12/2024.

    Bảng 19 cột, 4 hàng header, N hàng dữ liệu, 3 hàng footer.
    """
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
        from docx.shared import Pt, Cm, RGBColor
        from docx.oxml.ns import qn, nsdecls
        from docx.oxml import parse_xml
    except ImportError as exc:
        raise ImportError("Chưa cài python-docx. Chạy: pip install python-docx") from exc

    doc = Document()

    # ── Page setup: A4 ngang ──
    for section in doc.sections:
        section.page_width = Cm(29.7)
        section.page_height = Cm(21.0)
        section.left_margin = Cm(1.5)
        section.right_margin = Cm(1.5)
        section.top_margin = Cm(1.5)
        section.bottom_margin = Cm(1.5)

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)

    # ── Title ──
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(matrix.title or "MA TRẬN ĐỀ KIỂM TRA")
    run.bold = True
    run.font.size = Pt(13)

    if matrix.time_str:
        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r2 = p2.add_run(f"Thời gian: {matrix.time_str}")
        r2.font.size = Pt(11)

    doc.add_paragraph()

    # ── Analyze data ──
    groups = _group_rows_by_topic(matrix.rows)

    n_header = 4
    n_data = len(groups)
    n_footer = 3  # Tổng số câu, Tổng số điểm, Tỉ lệ %
    n_total_rows = n_header + n_data + n_footer

    table = doc.add_table(rows=n_total_rows, cols=_N_COLS)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # ── Helpers ──
    def _cell(ri, ci):
        return table.rows[ri].cells[ci]

    def _merge(r1, c1, r2, c2):
        _cell(r1, c1).merge(_cell(r2, c2))

    def _write(ri, ci, text, bold=False, center=True, size=10, valign=True):
        cell = _cell(ri, ci)
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = (WD_ALIGN_PARAGRAPH.CENTER if center
                       else WD_ALIGN_PARAGRAPH.LEFT)
        r = p.add_run(str(text))
        r.bold = bold
        r.font.size = Pt(size)
        r.font.name = "Times New Roman"
        if valign:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        return cell

    def _shade(ri, ci, color="D9E2F3"):
        """Tô nền cell."""
        tc = _cell(ri, ci)._tc
        shading = parse_xml(
            f'<w:shd {nsdecls("w")} w:fill="{color}" w:val="clear"/>')
        tc.get_or_add_tcPr().append(shading)

    # ══════════════════════════════════════════════════════════════
    #  4 HEADER ROWS
    # ══════════════════════════════════════════════════════════════

    # ── Cột 0: TT (merge row 0-3) ──
    _merge(0, 0, 3, 0)
    _write(0, 0, "TT", bold=True, size=10)

    # ── Cột 1: Chủ đề / Chương (merge row 0-3) ──
    _merge(0, 1, 3, 1)
    _write(0, 1, "Chủ đề /\nChương", bold=True, size=10)

    # ── Cột 2: Nội dung / Đơn vị kiến thức (merge row 0-3) ──
    _merge(0, 2, 3, 2)
    _write(0, 2, "Nội dung/\nĐơn vị\nkiến thức", bold=True, size=10)

    # ── Cột 3-14: Mức độ đánh giá (merge row 0, col 3-14) ──
    _merge(0, 3, 0, 14)
    _write(0, 3, "Mức độ đánh giá", bold=True, size=10)

    # ── Cột 15-17: Tổng (merge row 0-2, col 15-17) ──
    _merge(0, 15, 2, 17)
    _write(0, 15, "Tổng", bold=True, size=10)

    # ── Cột 18: Tỉ lệ % điểm số (merge row 0-3) ──
    _merge(0, 18, 3, 18)
    _write(0, 18, "Tỉ lệ %\ntổng\nđiểm", bold=True, size=10)

    # ── Row 1: TNKQ (col 3-11), Tự luận (col 12-14) ──
    _merge(1, 3, 1, 11)
    _write(1, 3, "Trắc nghiệm khách quan", bold=True, size=10)

    _merge(1, 12, 1, 14)
    _write(1, 12, "Tự luận", bold=True, size=10)

    # ── Row 2: Tên từng loại (mỗi loại merge 3 cột) ──
    for qt, start_c in _COL_QTYPE_START.items():
        _merge(2, start_c, 2, start_c + 2)
        _write(2, start_c, _QTYPE_LABEL[qt], bold=True, size=9)

    # ── Row 3: Nhận biết / Thông hiểu / Vận dụng (lặp 5 lần: 4 qtype + Tổng) ──
    for qt in _QTYPE_ORDER:
        sc = _COL_QTYPE_START[qt]
        for li, lv in enumerate(_LEVEL_ORDER):
            _write(3, sc + li, _LEVEL_LABEL[lv], bold=True, size=8)

    for li, lv in enumerate(_LEVEL_ORDER):
        _write(3, _COL_TOTAL_START + li, _LEVEL_LABEL[lv], bold=True, size=8)

    # ── Tô nền header ──
    header_color = "D9E2F3"
    for ri in range(4):
        for ci in range(_N_COLS):
            _shade(ri, ci, header_color)

    # ══════════════════════════════════════════════════════════════
    #  DATA ROWS (mỗi chủ đề = 1 hàng)
    # ══════════════════════════════════════════════════════════════

    # Tổng điểm toàn bộ (để tính %)
    total_points_all = sum(r.points for r in matrix.rows if r.enabled) or 1.0

    for gi, (topic, topic_rows) in enumerate(groups.items()):
        ri = n_header + gi

        _write(ri, 0, str(gi + 1), size=10)
        _write(ri, 1, topic, bold=False, center=False, size=10)

        # Nội dung kiến thức / Yêu cầu cần đạt
        req_text = _get_requirements(topic_rows)
        _write(ri, 2, req_text, bold=False, center=False, size=9)

        # Ô dữ liệu: số câu theo (qtype × level)
        for qt in _QTYPE_ORDER:
            sc = _COL_QTYPE_START[qt]
            for li, lv in enumerate(_LEVEL_ORDER):
                count = _count_questions(topic_rows, qt, lv)
                if count > 0:
                    _write(ri, sc + li, str(count), size=9)

        # Cột Tổng theo level
        for li, lv in enumerate(_LEVEL_ORDER):
            total_lv = _count_all_qtypes(topic_rows, lv)
            if total_lv > 0:
                _write(ri, _COL_TOTAL_START + li, str(total_lv), size=9)

        # Tỉ lệ %
        topic_pts = sum(r.points for r in topic_rows if r.enabled)
        pct = topic_pts / total_points_all * 100
        _write(ri, _COL_PERCENT, f"{pct:.0f}%", size=9)

    # ══════════════════════════════════════════════════════════════
    #  3 FOOTER ROWS
    # ══════════════════════════════════════════════════════════════

    footer_labels = ["Tổng số câu", "Tổng số điểm", "Tỉ lệ %"]
    footer_color = "E2EFDA"

    for fi, label in enumerate(footer_labels):
        fri = n_header + n_data + fi
        _merge(fri, 0, fri, 2)
        _write(fri, 0, label, bold=True, size=10)

        # Tô nền footer
        for ci in range(_N_COLS):
            _shade(fri, ci, footer_color)

        for qt in _QTYPE_ORDER:
            sc = _COL_QTYPE_START[qt]
            for li, lv in enumerate(_LEVEL_ORDER):
                if fi == 0:  # Tổng số câu
                    val = sum(_count_questions(rows, qt, lv)
                              for rows in groups.values())
                    if val > 0:
                        _write(fri, sc + li, str(val), bold=True, size=9)
                elif fi == 1:  # Tổng số điểm
                    val = sum(_sum_points(rows, qt, lv)
                              for rows in groups.values())
                    if val > 0:
                        txt = f"{val:g}"
                        _write(fri, sc + li, txt, bold=True, size=9)
                else:  # Tỉ lệ %
                    pts = sum(_sum_points(rows, qt, lv)
                              for rows in groups.values())
                    if pts > 0:
                        pct = pts / total_points_all * 100
                        _write(fri, sc + li, f"{pct:.0f}%", bold=True, size=9)

        # Cột Tổng theo level
        for li, lv in enumerate(_LEVEL_ORDER):
            if fi == 0:
                val = sum(_count_all_qtypes(rows, lv)
                          for rows in groups.values())
                if val > 0:
                    _write(fri, _COL_TOTAL_START + li, str(val),
                           bold=True, size=9)
            elif fi == 1:
                val = sum(_sum_points_all_qtypes(rows, lv)
                          for rows in groups.values())
                if val > 0:
                    _write(fri, _COL_TOTAL_START + li, f"{val:g}",
                           bold=True, size=9)
            else:
                pts = sum(_sum_points_all_qtypes(rows, lv)
                          for rows in groups.values())
                if pts > 0:
                    pct = pts / total_points_all * 100
                    _write(fri, _COL_TOTAL_START + li, f"{pct:.0f}%",
                           bold=True, size=9)

        # Cột Tỉ lệ % cuối
        if fi == 0:
            total_q = sum(1 for r in matrix.rows if r.enabled)
            _write(fri, _COL_PERCENT, str(total_q), bold=True, size=9)
        elif fi == 1:
            _write(fri, _COL_PERCENT, f"{total_points_all:g}",
                   bold=True, size=9)
        else:
            _write(fri, _COL_PERCENT, "100%", bold=True, size=9)

    # ══════════════════════════════════════════════════════════════
    #  COLUMN WIDTHS & FORMATTING
    # ══════════════════════════════════════════════════════════════
    try:
        table.columns[0].width = Cm(0.8)    # TT
        table.columns[1].width = Cm(3.0)    # Chủ đề
        table.columns[2].width = Cm(4.0)    # Nội dung
        for ci in range(3, 18):
            table.columns[ci].width = Cm(1.4)
        table.columns[18].width = Cm(1.5)   # Tỉ lệ %
    except Exception:
        pass

    # Cell padding
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(1)

    doc.save(str(Path(output_path)))
