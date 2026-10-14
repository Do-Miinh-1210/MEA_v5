"""
export/report_exporter.py — Xuất báo cáo lỗi câu hỏi ra Excel/CSV (Yêu cầu 1).
"""
from __future__ import annotations
from pathlib import Path
from typing import Optional
import csv, datetime

from core.models import Question, ExamData
from config import TYPE_LABELS


def _q_summary(q: Question) -> str:
    if q.qtype == "TN":
        return f"{q.question[:80]} | Đáp án: {q.answer}"
    elif q.qtype == "DS":
        items = " / ".join(f"({k}){v[:30]}" for k,v in q.items.items())
        return f"{q.question[:60]} | {items}"
    return q.question[:100]


def export_error_report_csv(exam: ExamData, path: str) -> int:
    """Xuất báo cáo lỗi ra CSV. Trả về số dòng lỗi."""
    rows = []
    for i, q in enumerate(exam.questions, 1):
        has_err = q._error
        review  = q._review or {}
        ok      = review.get("ok", True)
        if not has_err and ok:
            continue
        rows.append({
            "STT": i,
            "Loại": TYPE_LABELS.get(q.qtype, q.qtype),
            "Mức": q.level,
            "Chủ đề": q.topic,
            "Tóm tắt câu hỏi": _q_summary(q),
            "Lỗi tạo": "Có" if has_err else "",
            "Lỗi review": "" if ok else review.get("issue",""),
            "Gợi ý sửa": review.get("tip",""),
            "Điểm review": review.get("score","") if review else "",
            "Thời gian": datetime.datetime.now().strftime("%Y-%m-%d %H:%M"),
        })

    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        if not rows:
            f.write("Không có lỗi nào được phát hiện.\n")
            return 0
        writer = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)
    return len(rows)


def export_error_report_excel(exam: ExamData, path: str) -> int:
    """Xuất báo cáo lỗi ra Excel (cần openpyxl). Trả về số dòng lỗi."""
    try:
        import openpyxl
        from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    except ImportError:
        raise ImportError("Chưa cài openpyxl\nChạy: pip install openpyxl")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Báo cáo lỗi"

    RED_FILL  = PatternFill("solid", fgColor="FFCCCC")
    YLW_FILL  = PatternFill("solid", fgColor="FFF2CC")
    HDR_FILL  = PatternFill("solid", fgColor="1F497D")
    HDR_FONT  = Font(bold=True, color="FFFFFF")
    CENTER    = Alignment(horizontal="center", vertical="center", wrap_text=True)
    LEFT      = Alignment(horizontal="left",   vertical="top",    wrap_text=True)
    thin      = Side(style="thin", color="AAAAAA")
    BORDER    = Border(left=thin, right=thin, top=thin, bottom=thin)

    headers = ["STT","Loại","Mức","Chủ đề","Tóm tắt câu hỏi",
               "Lỗi tạo","Lỗi review","Gợi ý sửa","Điểm review"]
    widths  = [6,    16,    8,    22,    50,          10,     30,          35,         12]

    for col, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill   = HDR_FILL
        cell.font   = HDR_FONT
        cell.alignment = CENTER
        cell.border = BORDER
        ws.column_dimensions[cell.column_letter].width = w

    ws.row_dimensions[1].height = 22

    data_row = 2
    err_count = 0
    for i, q in enumerate(exam.questions, 1):
        has_err = q._error
        review  = q._review or {}
        ok      = review.get("ok", True)
        if not has_err and ok:
            continue

        fill = RED_FILL if has_err else YLW_FILL
        vals = [
            i,
            TYPE_LABELS.get(q.qtype, q.qtype),
            q.level,
            q.topic,
            _q_summary(q),
            "⚠️ Lỗi" if has_err else "",
            review.get("issue","") if not ok else "",
            review.get("tip",""),
            review.get("score","") if review else "",
        ]
        for col, val in enumerate(vals, 1):
            cell = ws.cell(row=data_row, column=col, value=val)
            cell.fill      = fill
            cell.alignment = CENTER if col <= 3 else LEFT
            cell.border    = BORDER

        ws.row_dimensions[data_row].height = 40
        data_row += 1
        err_count += 1

    if err_count == 0:
        ws.cell(row=2, column=1, value="✅ Không phát hiện lỗi nào.")

    # Summary sheet
    ws2 = wb.create_sheet("Tổng quan")
    ws2.append(["Môn", exam.subject])
    ws2.append(["Lớp", exam.grade])
    ws2.append(["Tổng câu", len(exam.questions)])
    ws2.append(["Câu lỗi tạo", sum(1 for q in exam.questions if q._error)])
    ws2.append(["Câu lỗi review", sum(1 for q in exam.questions
                                      if q._review and not q._review.get("ok",True))])
    ws2.append(["Xuất lúc", datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")])

    wb.save(path)
    return err_count
